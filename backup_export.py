# backup_export.py — 予約全件・顧客集計・POS売上明細をまとめたExcelを作る共通ロジック
# app.py の /backup/xlsx（手動ダウンロード）と daily_backup.py（毎日自動メール送信）の
# 両方から使われる。出力内容がズレないよう、ここに一本化している。
import io
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from database import get_all_reservations, get_pos_sales
from logic import aggregate_customer_ranking

JST = timezone(timedelta(hours=9))


def today_jst() -> date:
    return datetime.now(JST).date()


def build_backup_workbook_bytes() -> bytes:
    all_res   = get_all_reservations(include_cancelled=True)
    customers = aggregate_customer_ranking([r for r in all_res if r['status'] != 'cancelled'])
    customers.sort(key=lambda c: -c['visit_count'])
    pos_rows  = get_pos_sales()

    wb = Workbook()

    ws1 = wb.active
    ws1.title = '予約全件'
    ws1.append([
        'ID', '日付', '時間', '氏名', '電話番号',
        '大人', '子供', '合計人数', '利用時間(分)',
        '個室', 'VIP', '団体', '予算/人', 'ニーズ',
        '男性', '女性', '幹事メモ', '備考', 'メニュー', '売上',
        'テーブル', 'ステータス', '来店受付日時', '登録日時', '更新日時',
    ])
    for r in all_res:
        ws1.append([
            r['id'], r['date'], r['time_slot'], r['name'], r['phone'],
            r['adults'], len(r['children_info']), r['total_people'], r['duration_minutes'],
            '○' if r['private_room'] else '', '○' if r['is_vip'] else '', '○' if r['is_group'] else '',
            r.get('budget_per_person') or '', r.get('needs_type') or '',
            r.get('gender_male') or '', r.get('gender_female') or '',
            r.get('organizer_note') or '', r.get('notes') or '',
            r.get('menu_note') or '', r.get('sales_amount') or '',
            '・'.join(str(t) for t in r['assigned_tables']), r['status'], r.get('visited_at') or '',
            r['created_at'], r['updated_at'],
        ])

    ws2 = wb.create_sheet('顧客集計')
    ws2.append(['氏名', '電話番号', '来店回数', '累計売上', '平均単価/回',
                '初回来店', '前回来店', '前回からの経過日数', '最終来店受付日時'])
    for c in customers:
        ws2.append([
            c['name'], c['phone'], c['visit_count'], c['total_sales'], c['avg_sales'],
            c['first_visit_date'], c['last_visit_date'], c['days_since_last_visit'],
            c.get('last_visited_at') or '',
        ])

    ws3 = wb.create_sheet('POS売上明細')
    ws3.append(['日付', '時刻', '伝票番号', '卓番号', '客数', '商品名', '金額',
                'お客様名', '電話番号', '突合ステータス', '紐付け予約ID', 'インポート回次'])
    for r in pos_rows:
        ws3.append([
            r['sale_date'], r['sale_time'] or '', r['receipt_no'] or '', r['table_no'] or '',
            r['party_size'] or '', r['item_name'] or '', r['amount'],
            r['customer_name'] or '', r['phone'] or '',
            r['match_status'], r['matched_reservation_id'] or '', r['import_batch'],
        ])

    for ws in (ws1, ws2, ws3):
        for col_cells in ws.columns:
            length = max((len(str(cell.value)) for cell in col_cells if cell.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
